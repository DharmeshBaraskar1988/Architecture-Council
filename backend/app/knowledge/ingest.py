"""PRD ingestion: read PDF, Word, HTML, Markdown or text into pages, then into page-aware chunks.

Every chunk keeps its page number so any fact can be traced back to where it was written.
Markdown and text have no real pages: mark them with lines like `<!-- page: 2 -->`
(or form-feed characters). Without markers, each top-level heading (# or ##) starts a new
section, numbered like pages; without headings the whole file is page 1.
HTML is split the same way at <h1>/<h2>. Each page or section keeps its title for the outline.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional

_PAGE_MARK = re.compile(r"^\s*<!--\s*page:\s*(\d+)\s*-->\s*$", re.I | re.M)
_HEADING = re.compile(r"^#{1,2}\s+(.+?)\s*#*\s*$", re.M)
_ANY_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*#*\s*$", re.M)


@dataclass
class Page:
    number: int
    text: str
    title: Optional[str] = None      # first heading on the page (or first line), for the document outline

    def __post_init__(self):
        if self.title is None:
            m = _ANY_HEADING.search(self.text)
            line = m.group(1) if m else next((ln.strip() for ln in self.text.splitlines() if ln.strip()), "")
            self.title = line[:80] or None


@dataclass
class Chunk:
    id: str
    page: int
    text: str


@dataclass
class Document:
    name: str
    pages: list[Page] = field(default_factory=list)

    def page_text(self, n: int) -> str:
        for p in self.pages:
            if p.number == n:
                return p.text
        return ""

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.pages)


def _split_marked(text: str) -> list[Page]:
    marks = list(_PAGE_MARK.finditer(text))
    if marks:
        pages = []
        head = text[:marks[0].start()].strip()   # title above the first marker joins the first page
        for i, m in enumerate(marks):
            end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
            body = text[m.end():end].strip()
            if i == 0 and head:
                body = head + "\n\n" + body
            pages.append(Page(int(m.group(1)), body))
        return pages
    if "\f" in text:
        return [Page(i + 1, t.strip()) for i, t in enumerate(text.split("\f"))]
    return _split_headings(text)


def _split_headings(text: str) -> list[Page]:
    """No page markers: every # or ## heading starts a section (text above the first joins it)."""
    heads = list(_HEADING.finditer(text))
    if len(heads) < 2:
        return [Page(1, text.strip())]
    starts = [0] + [h.start() for h in heads[1:]]
    ends = starts[1:] + [len(text)]
    parts = [text[a:b].strip() for a, b in zip(starts, ends)]
    return [Page(i + 1, t) for i, t in enumerate(p for p in parts if p)]


class _HtmlText(HTMLParser):
    """Visible text of an HTML page, with headings written as Markdown so sections split the same way."""
    SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    BLOCK = {"p", "div", "section", "article", "li", "tr", "br", "table", "ul", "ol", "header", "footer",
             "main", "blockquote", "pre"}
    HEADS = {"h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.HEADS:
            self.out.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag in self.BLOCK:
            self.out.append("\n")
        elif tag in ("td", "th"):
            self.out.append(" | ")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
        elif tag in self.HEADS or tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> str:
    p = _HtmlText()
    p.feed(html)
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(p.out))
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def load_document(path: str | Path) -> Document:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".md", ".markdown", ".txt"):
        return Document(path.name, _split_marked(path.read_text(encoding="utf-8")))
    if ext in (".html", ".htm"):
        return Document(path.name, _split_headings(html_to_text(path.read_text(encoding="utf-8", errors="replace"))))
    if ext == ".pdf":
        from pypdf import PdfReader  # BSD-3-Clause
        reader = PdfReader(str(path))
        return Document(path.name, [Page(i + 1, (pg.extract_text() or "").strip()) for i, pg in enumerate(reader.pages)])
    if ext == ".docx":
        import docx  # python-docx, MIT
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = docx.Document(str(path))
        # Word has no fixed pages: use manual page breaks and the page breaks Word last rendered
        # (approximate). Without either, headings split the document into sections, as for Markdown.
        # Paragraphs and tables are read in document order, so a table stays in its section.
        pages, buf, n, breaks = [], [], 1, False
        for el in d.element.body.iterchildren():
            if el.tag.endswith("}p"):
                para = Paragraph(el, d)
                xml = el.xml
                if ('w:br w:type="page"' in xml or "<w:lastRenderedPageBreak/>" in xml) and buf:
                    pages.append(Page(n, "\n".join(buf).strip()))
                    buf, n, breaks = [], n + 1, True
                style = (para.style.name if para.style is not None else "") or ""
                level = int(style.split()[-1]) if style.startswith("Heading") and style.split()[-1].isdigit() else \
                    1 if style == "Title" else 0
                buf.append(("#" * level + " " if level and para.text.strip() else "") + para.text)
            elif el.tag.endswith("}tbl"):
                for row in Table(el, d).rows:
                    buf.append(" | ".join(c.text.strip() for c in row.cells))
        if breaks:
            pages.append(Page(n, "\n".join(buf).strip()))
            return Document(path.name, [p for p in pages if p.text])
        return Document(path.name, _split_headings("\n".join(buf)))
    raise ValueError(f"unsupported file type: {ext} (use .pdf, .docx, .html, .md or .txt)")


def chunk(doc: Document, max_chars: int = 2500, overlap: int = 200) -> list[Chunk]:
    """Split each page into overlapping chunks; chunks never cross pages."""
    out: list[Chunk] = []
    for p in doc.pages:
        t, i, k = p.text, 0, 0
        while i < len(t):
            piece = t[i:i + max_chars]
            out.append(Chunk(id=f"p{p.number}-c{k}", page=p.number, text=piece))
            k += 1
            if i + max_chars >= len(t):
                break
            i += max_chars - overlap
    return out
